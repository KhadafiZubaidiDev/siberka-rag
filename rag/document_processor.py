"""
rag/document_processor.py — Load & chunk dokumen
==================================================
Mendukung PDF, TXT, dan DOCX.
"""
from __future__ import annotations

import hashlib
import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

import config

logger = logging.getLogger(__name__)


# ── Data Classes ─────────────────────────────────────────────────

@dataclass
class ProcessedDocument:
    """Hasil pemrosesan satu file."""
    filename: str
    file_hash: str
    chunks: list[Document]
    total_pages: int = 0
    total_chars: int = 0
    extra_meta: dict = field(default_factory=dict)


# ── Text Splitter ─────────────────────────────────────────────────

def _make_splitter() -> RecursiveCharacterTextSplitter:
    return RecursiveCharacterTextSplitter(
        chunk_size=config.CHUNK_SIZE,
        chunk_overlap=config.CHUNK_OVERLAP,
        length_function=len,
        separators=["\n\n", "\n", ".", "!", "?", ",", " ", ""],
    )


# ── Loaders ───────────────────────────────────────────────────────

def _load_pdf(file_bytes: bytes, filename: str) -> tuple[list[Document], int]:
    """Muat PDF menggunakan pypdf."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    docs: list[Document] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(Document(
                page_content=text,
                metadata={"source": filename, "page": i + 1},
            ))
    return docs, len(reader.pages)


def _load_txt(file_bytes: bytes, filename: str) -> tuple[list[Document], int]:
    """Muat file teks biasa."""
    text = file_bytes.decode("utf-8", errors="replace")
    doc = Document(page_content=text, metadata={"source": filename, "page": 1})
    return [doc], 1


def _load_docx(file_bytes: bytes, filename: str) -> tuple[list[Document], int]:
    """Muat DOCX menggunakan python-docx."""
    from docx import Document as DocxDoc

    docx = DocxDoc(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in docx.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    doc = Document(page_content=text, metadata={"source": filename, "page": 1})
    return [doc], len(paragraphs)


# ── Public API ────────────────────────────────────────────────────

def process_uploaded_file(
    file_bytes: bytes,
    filename: str,
) -> ProcessedDocument:
    """
    Proses file yang diunggah menjadi chunks siap di-embed.

    Args:
        file_bytes: Konten file dalam bytes.
        filename:   Nama file asli (digunakan untuk deteksi ekstensi).

    Returns:
        ProcessedDocument dengan list chunks (LangChain Documents).

    Raises:
        ValueError: Jika ekstensi file tidak didukung.
    """
    ext = Path(filename).suffix.lower()
    if ext not in config.SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Format '{ext}' tidak didukung. "
            f"Gunakan: {', '.join(config.SUPPORTED_EXTENSIONS)}"
        )

    # Hash untuk deduplikasi
    file_hash = hashlib.sha256(file_bytes).hexdigest()

    # Load berdasarkan tipe
    if ext == ".pdf":
        raw_docs, total_pages = _load_pdf(file_bytes, filename)
    elif ext == ".txt":
        raw_docs, total_pages = _load_txt(file_bytes, filename)
    elif ext == ".docx":
        raw_docs, total_pages = _load_docx(file_bytes, filename)
    else:
        raise ValueError(f"Ekstensi tidak dikenal: {ext}")

    total_chars = sum(len(d.page_content) for d in raw_docs)

    # Splitting
    splitter = _make_splitter()
    chunks = splitter.split_documents(raw_docs)

    # Tambahkan metadata chunk
    for idx, chunk in enumerate(chunks):
        chunk.metadata.update({
            "chunk_index": idx,
            "file_hash": file_hash,
            "total_chunks": len(chunks),
        })

    logger.info(
        "Processed '%s': %d pages, %d chars, %d chunks",
        filename, total_pages, total_chars, len(chunks),
    )

    return ProcessedDocument(
        filename=filename,
        file_hash=file_hash,
        chunks=chunks,
        total_pages=total_pages,
        total_chars=total_chars,
    )
